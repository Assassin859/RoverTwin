"""Fill TECHFEST Round-1 template with SatTwin ST-09 content.

Source template: Downloads/TECHFEST-2026-27-Round-1-Template-Clean.pptx
Outputs: docs/SatTwin-ST09.pptx and docs/RoverTwin-ST09.pptx
"""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path.home() / "Downloads" / "TECHFEST-2026-27-Round-1-Template-Clean.pptx"
OUT = ROOT / "docs" / "SatTwin-ST09.pptx"
OUT_ALIAS = ROOT / "docs" / "RoverTwin-ST09.pptx"

CYAN = RGBColor(0x22, 0xD3, 0xEE)
SKY = RGBColor(0x7D, 0xD3, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xBA, 0xE6, 0xFD)
SAFFRON = RGBColor(0xFF, 0x99, 0x33)


def _set_run(run, text, *, size=16, bold=False, color=WHITE, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color


def add_body(slide, lines, *, top_in=1.35, left_in=0.6, width_in=12.1, size=16, gap=8):
    """Content block under the template section title."""
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(5.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.clear()
        run = p.add_run()
        is_head = line.startswith("▸") or line.endswith(":") and len(line) < 40
        col = SAFFRON if line.startswith("▸") else (SKY if is_head else WHITE)
        _set_run(run, line, size=size if not line.startswith("  ") else size - 1, bold=is_head or line.startswith("▸"), color=col, font="Calibri")
        p.space_after = Pt(gap)
        p.level = 0


def add_project_banner(slide):
    """Project identity on the cover (between hackathon title and ROUND 1 band)."""
    box = slide.shapes.add_textbox(Inches(0.6), Inches(3.15), Inches(12.1), Inches(1.35))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "SatTwin  ·  ST-09", size=32, bold=True, color=WHITE, font="Bahnschrift")
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    _set_run(r2, "Mission Digital Twin for Predictive Fault Simulation", size=18, color=MUTED, font="Calibri")
    p3 = tf.add_paragraph()
    r3 = p3.add_run()
    _set_run(r3, "LEO EO smallsat twin  ·  github.com/Assassin859/RoverTwin  ·  rovertwin.vercel.app", size=13, color=CYAN, font="Consolas")


def main():
    if not TEMPLATE.is_file():
        raise SystemExit(f"missing template: {TEMPLATE}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE, OUT)
    prs = Presentation(str(OUT))

    # 1 Cover — keep TECHFEST chrome; add SatTwin banner
    add_project_banner(prs.slides[0])

    # 2 Challenge / Problem Statement
    add_body(prs.slides[1], [
        "▸ TECHFEST ST-09 — Mission Digital Twin for Predictive Fault Simulation",
        "",
        "Build a digital twin that is not a dashboard of canned plots:",
        "  • Model ≥3 coupled spacecraft subsystems with real cause→effect links",
        "  • Synchronise to telemetry (or honestly show pass-gated sync states)",
        "  • Inject faults, show where effects spread next, and why (equations)",
        "  • Predict impact with time + uncertainty; rank recovery before uplink",
        "",
        "Judge traps: pretty UI with no coupling · standalone sim · no validation story",
    ], size=17)

    # 3 Mission analysis / problem understanding
    add_body(prs.slides[2], [
        "▸ Domain — 500 km LEO Earth-observation smallsat (~95 min orbit)",
        "  Eclipse + ground-station passes gate TM/TC. Ops must act between AOS windows.",
        "",
        "▸ Failure physics is multi-hop",
        "  Battery degradation → heat (EPS→TCS) → ADCS gyro bias / COMMS margin / PAYLOAD shed",
        "  Operators need root vs knock-on, not six independent gauges.",
        "",
        "▸ What “twin” means here",
        "  Same physics model as the craft, corrected only by downlinked frames.",
        "  Sync states INIT / SYNCED / LOW RATE / BLIND — silence is evidence.",
    ], size=16)

    # 4 Proposed solution
    add_body(prs.slides[3], [
        "▸ SatTwin — live mission digital twin + operator console",
        "",
        "  1. Mirror — six coupled subsystems (EPS · TCS · ADCS · COMMS · PAYLOAD · OBDH)",
        "  2. Break — inject battery / thermal / sensor / comms faults",
        "  3. Spread — 14 equation-backed coupling edges light up as the cascade runs",
        "  4. Predict — 2 h ensemble forecast + first-critical headline with band",
        "  5. Recover — ranked plans scored for safety / power / thermal / data / pointing",
        "  6. Execute — queue for next AOS; confirm uplink; show contained / resolved",
        "",
        "Hands-free 3-min judge path on live backend or Vercel offline replay.",
    ], size=16)

    # 5 System architecture
    add_body(prs.slides[4], [
        "▸ Data path (enforced by tests/test_twin.py)",
        "  Plant (hidden truth) → RadioLink (latency, rate limit, loss, AOS/LOS)",
        "                      → DigitalTwin (same model.step + estimated health only)",
        "",
        "▸ Ground segment",
        "  Correlate / cascade graph · Predictor (ensemble + plan simulation)",
        "  FastAPI: private WebSocket /ws (~10 Hz snaps) + REST /api/* + SQLite",
        "  View-model: incident · ops · timeline · score_breakdown (single source of truth)",
        "",
        "▸ Operator console",
        "  Inject · cascade · forecast · plans · evidence (/api/validate, /api/backtest, CSV)",
    ], size=15)

    # 6 Technology stack
    add_body(prs.slides[5], [
        "▸ Backend",
        "  Python · FastAPI · WebSocket · physics model (1 Hz) · pytest suite",
        "  Optional local Ollama (qwen2.5:3b) — narrates twin facts only",
        "",
        "▸ Frontend",
        "  Vanilla ops console (HTML/CSS/JS) · Three.js orbit scene · uPlot-style charts",
        "  Offline tape: demo-replay.json.gz on Vercel when backend unreachable",
        "",
        "▸ Delivery",
        "  Local uvicorn for venue demo · https://rovertwin.vercel.app backup",
        "  ?backend= points static UI at a live ground segment",
    ], size=16)

    # 7 Workflow
    add_body(prs.slides[6], [
        "▸ 3-minute judge demo",
        "  Landing → 3-MIN JUDGE PATH → inject battery → watch EPS→TCS→… cascade",
        "  Forecast pauses on first critical → compare / execute isolate at AOS",
        "  Confirm recovery · optional Show truth (harness) · hover edges for KaTeX/equations",
        "",
        "▸ Operator loop (always)",
        "  Observe sync + subsystem scores → diagnose root → simulate plans → uplink at pass",
        "",
        "Cold start:  uvicorn backend.app:app --port 8000",
    ], size=16)

    # 8 Impact
    add_body(prs.slides[7], [
        "▸ Innovation",
        "  Cause→effect is model-derived, not scripted UI animation",
        "  Twin never cheats by reading plant state — same constraint as a real ground segment",
        "  Plans are forward-simulated and scored before the next uplink opportunity",
        "",
        "▸ Impact for ST-09",
        "  Answers all four judge questions on-screen (no slide-only story)",
        "  Venue-ready: live laptop demo + offline Vercel replay fallback",
        "  Extensible: CSV ingest, validation/backtest APIs, optional LLM narration",
    ], size=16)

    # 9 Roadmap
    add_body(prs.slides[8], [
        "▸ Now (Round 1)",
        "  LEO smallsat twin · 6 subsystems · cascade · forecast · recovery · evidence APIs",
        "",
        "▸ Next",
        "  Richer ephemeris / multi-GS · flight-correlated ensembles · more fault catalogues",
        "  Hardened React console (optional) behind the same protocol v2 contract",
        "",
        "▸ Future scope",
        "  Ops training scenarios · hardware-in-the-loop telemetry adapters · multi-sat constellation",
    ], size=16)

    # 10 Team
    add_body(prs.slides[9], [
        "▸ Team SatTwin",
        "",
        "  Project: SatTwin (repo Assassin859/RoverTwin)",
        "  Track: Space Technology Hackathon · Problem ST-09",
        "  Demo: local uvicorn + https://rovertwin.vercel.app",
        "",
        "  Add member names / roles here before venue submit.",
    ], size=18)

    # 11 Thank you — leave template text; add one contact line
    box = prs.slides[10].shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(12.1), Inches(0.8))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "SatTwin · ST-09  ·  github.com/Assassin859/RoverTwin", size=16, color=CYAN, font="Consolas")
    p.alignment = PP_ALIGN.CENTER

    prs.save(OUT)
    shutil.copy2(OUT, OUT_ALIAS)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print(f"wrote {OUT_ALIAS}")


if __name__ == "__main__":
    main()
